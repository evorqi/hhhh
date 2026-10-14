"""CSV and XLSX readers. Declared worksheet dimensions are not trusted blindly."""

from __future__ import annotations

import csv
import io
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import Cell as OpenpyxlCell
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from cleaning_workspace.config import Settings
from cleaning_workspace.domain.models import Row, Sheet, WorkbookModel, make_cell
from cleaning_workspace.domain.security import Detection, warning
from cleaning_workspace.errors import AppError


def _looks_textual(text: str) -> bool:
    if not text:
        return False
    controls = sum(1 for char in text if ord(char) < 32 and char not in "\r\n\t")
    if controls > max(1, len(text) * 0.02):
        return False
    printable = sum(1 for char in text if char.isprintable() or char in "\r\n\t")
    return printable / len(text) > 0.95


def decode_csv_bytes(data: bytes) -> tuple[str, str, bool]:
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig"), "utf-8-sig", True
    try:
        return data.decode("utf-8"), "utf-8", False
    except UnicodeDecodeError:
        pass
    text = data.decode("cp1256")
    if _looks_textual(text):
        return text, "cp1256", False
    raise AppError(
        "The text encoding is not supported. Save the CSV as UTF-8.",
        "رمزگذاری متن پشتیبانی نمی‌شود. فایل CSV را با UTF-8 ذخیره کنید.",
        code="unsupported_encoding",
        status_code=422,
    )


def detect_delimiter(text: str) -> str:
    sample = text[:50_000]
    best = ","
    best_score = -1
    for delimiter in (",", ";", "\t", "|"):
        rows = list(csv.reader(io.StringIO(sample), delimiter=delimiter))
        useful = [row for row in rows[:40] if any(cell.strip() for cell in row)]
        if not useful:
            continue
        widths = [len(row) for row in useful]
        score = sum(widths)
        if len(set(widths)) == 1 and max(widths) > 1:
            score += 100
        elif max(widths) > 1:
            score += 10
        if score > best_score:
            best_score = score
            best = delimiter
    return best


def read_model(path: Path, detection: Detection, settings: Settings) -> WorkbookModel:
    data = path.read_bytes()
    if detection.source_format == "csv":
        return _read_csv(data, detection, settings)
    return _read_xlsx(data, detection, settings)


def _read_csv(data: bytes, detection: Detection, settings: Settings) -> WorkbookModel:
    text, encoding, had_bom = decode_csv_bytes(data)
    delimiter = detect_delimiter(text)
    parsed = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if len(parsed) > settings.max_rows:
        raise AppError(
            "The CSV exceeds the row limit.",
            "تعداد ردیف‌های CSV از حد مجاز بیشتر است.",
            code="too_many_rows",
            status_code=422,
        )
    width = max((len(row) for row in parsed), default=0)
    if width > settings.max_cols:
        raise AppError(
            "The CSV exceeds the column limit.",
            "تعداد ستون‌های CSV از حد مجاز بیشتر است.",
            code="too_many_columns",
            status_code=422,
        )
    if len(parsed) * max(width, 1) > settings.max_cells:
        raise AppError(
            "The CSV exceeds the cell limit.",
            "تعداد سلول‌های CSV از حد مجاز بیشتر است.",
            code="too_many_cells",
            status_code=422,
        )
    warnings = list(detection.warnings)
    if encoding == "cp1256":
        warnings.append(
            warning(
                "decoded_as_cp1256",
                "UTF-8 decoding failed. The file was decoded as Windows-1256.",
                "خواندن UTF-8 ممکن نشد. فایل با رمز Windows-1256 خوانده شد.",
            )
        )
    widths = {len(row) for row in parsed if any(cell.strip() for cell in row)}
    if len(widths) > 1:
        warnings.append(
            warning(
                "uneven_rows",
                "CSV rows do not all have the same number of fields.",
                "طول ردیف‌های CSV یکسان نیست.",
            )
        )
    rows: list[Row] = []
    for index, raw in enumerate(parsed, start=1):
        cells = {}
        for column, value in enumerate(raw, start=1):
            if value == "":
                continue
            cells[column] = make_cell(value)
        rows.append(Row(original_index=index, cells=cells))
    actual_max_row = max((row.original_index for row in rows if row.cells), default=0)
    sheet = Sheet(
        name="CSV",
        rows=rows,
        actual_max_row=actual_max_row,
        actual_max_col=width,
        declared_max_row=len(parsed),
        declared_max_col=width,
    )
    return WorkbookModel(
        source_format="csv",
        sheets=[sheet],
        delimiter=delimiter,
        encoding=encoding,
        had_bom=had_bom,
        extension_mismatch=detection.extension_mismatch,
        warnings=warnings,
    )


def _read_xlsx(data: bytes, detection: Detection, settings: Settings) -> WorkbookModel:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_vba=detection.has_vba)
    except Exception as exc:
        raise AppError(
            "The spreadsheet could not be opened.",
            "کارپوشه باز نشد. فایل ممکن است خراب باشد.",
            code="corrupt_xlsx",
            status_code=422,
        ) from exc
    if len(workbook.sheetnames) > settings.max_sheets:
        raise AppError(
            "The workbook has too many sheets.",
            "تعداد برگه‌ها از حد مجاز بیشتر است.",
            code="too_many_sheets",
            status_code=422,
        )
    warnings = list(detection.warnings)
    sheets: list[Sheet] = []
    total_cells = 0
    for worksheet in workbook.worksheets:
        sheet, cell_count = _read_worksheet(worksheet, settings, warnings)
        total_cells += cell_count
        if total_cells > settings.max_cells:
            raise AppError(
                "The workbook exceeds the cell limit.",
                "تعداد سلول‌های کارپوشه از حد مجاز بیشتر است.",
                code="too_many_cells",
                status_code=422,
            )
        sheets.append(sheet)
    defined_names = 0
    try:
        defined_names = len(list(workbook.defined_names.keys()))
    except Exception:
        defined_names = 0
    external_links = len(getattr(workbook, "_external_links", []) or [])
    if external_links:
        warnings.append(
            warning(
                "external_links",
                "The workbook has external links. They are not fetched or executed.",
                "کارپوشه پیوند بیرونی دارد. این پیوندها دریافت یا اجرا نمی‌شوند.",
            )
        )
    epoch = getattr(workbook, "epoch", None)
    date1904 = epoch == CALENDAR_MAC_1904 or getattr(epoch, "year", None) == 1904
    if date1904:
        warnings.append(
            warning(
                "date_system_1904",
                "This workbook uses the 1904 date system. Excel date values are read through the library and are not recalculated.",
                "این کارپوشه از سامانه تاریخ ۱۹۰۴ استفاده می‌کند. تاریخ‌های عددی اکسل فقط از راه کتابخانه خوانده می‌شوند و دوباره محاسبه نمی‌شوند.",
            )
        )
    return WorkbookModel(
        source_format="xlsx",
        sheets=sheets,
        date1904=date1904,
        has_vba=detection.has_vba,
        defined_name_count=defined_names,
        external_link_count=external_links,
        extension_mismatch=detection.extension_mismatch,
        warnings=warnings,
    )


def _read_worksheet(worksheet, settings: Settings, warnings: list[dict[str, str]]) -> tuple[Sheet, int]:
    occupied: dict[tuple[int, int], OpenpyxlCell] = {}
    comments = 0
    hyperlinks = 0
    formulas = 0
    for key, cell in worksheet._cells.items():
        if isinstance(cell, MergedCell) or not isinstance(cell, OpenpyxlCell):
            continue
        if cell.value is None and cell.data_type != "f":
            continue
        occupied[key] = cell
        if cell.comment is not None:
            comments += 1
        if cell.hyperlink is not None:
            hyperlinks += 1
        if cell.data_type == "f":
            formulas += 1
    if len(occupied) > settings.max_cells:
        raise AppError(
            "A worksheet exceeds the cell limit.",
            "یکی از برگه‌ها از حد سلول مجاز بیشتر است.",
            code="too_many_cells",
            status_code=422,
        )
    actual_max_row = max((row for row, _column in occupied), default=0)
    actual_max_col = max((column for _row, column in occupied), default=0)
    for merged in worksheet.merged_cells.ranges:
        bounds = str(merged).split(":")
        if len(bounds) != 2:
            continue
        digits = "".join(char for char in bounds[1] if char.isdigit())
        if not digits:
            continue
        merge_end = int(digits)
        if merge_end <= actual_max_row:
            continue
        if merge_end - actual_max_row > 500 or merge_end > settings.max_rows:
            warnings.append(
                warning(
                    "large_merge",
                    f"Sheet {worksheet.title} has a merged range ending at row {merge_end}, beyond the cells that were read.",
                    f"برگه {worksheet.title} سلول ادغام‌شده‌ای تا ردیف {merge_end} دارد که خارج از سلول‌های خوانده‌شده است.",
                )
            )
            continue
        actual_max_row = merge_end
    if actual_max_row > settings.max_rows or actual_max_col > settings.max_cols:
        raise AppError(
            "A worksheet exceeds the row or column limit.",
            "یکی از برگه‌ها از حد ردیف یا ستون مجاز بیشتر است.",
            code="too_many_rows",
            status_code=422,
        )
    declared_row = int(worksheet.max_row or 0)
    declared_col = int(worksheet.max_column or 0)
    if declared_row > actual_max_row + 10 or declared_col > actual_max_col + 10:
        warnings.append(
            warning(
                "inflated_dimension",
                f"Sheet {worksheet.title} declares a used range larger than the cells that contain values. Empty declared cells were not materialized.",
                f"برگه {worksheet.title} محدوده‌ای بزرگ‌تر از سلول‌های دارای مقدار اعلام کرده است. سلول‌های خالی اعلام‌شده ساخته نشدند.",
            )
        )
    rows: list[Row] = []
    for row_index in range(1, actual_max_row + 1):
        cells = {}
        for column in range(1, actual_max_col + 1):
            cell = occupied.get((row_index, column))
            if cell is None:
                continue
            if cell.data_type == "f":
                cells[column] = make_cell(
                    formula=str(cell.value) if cell.value is not None else "=",
                    number_format=cell.number_format,
                )
            else:
                cells[column] = make_cell(cell.value, number_format=cell.number_format)
        rows.append(Row(original_index=row_index, cells=cells))
    hidden_rows = 0
    for row_index in range(1, actual_max_row + 1):
        dimension = worksheet.row_dimensions.get(row_index) if hasattr(worksheet, "row_dimensions") else None
        if dimension is not None and getattr(dimension, "hidden", False):
            hidden_rows += 1
    hidden_cols = 0
    for column in range(1, actual_max_col + 1):
        letter = openpyxl.utils.get_column_letter(column)
        dimension = worksheet.column_dimensions.get(letter)
        if dimension is not None and getattr(dimension, "hidden", False):
            hidden_cols += 1
    merged = [str(item) for item in worksheet.merged_cells.ranges]
    charts = len(getattr(worksheet, "_charts", []) or [])
    images = len(getattr(worksheet, "_images", []) or [])
    validations = 0
    if worksheet.data_validations is not None:
        validations = len(worksheet.data_validations.dataValidation)
    tables = len(getattr(worksheet, "tables", {}) or {})
    if charts or images:
        warnings.append(
            warning(
                "objects_preserved_with_limits",
                f"Sheet {worksheet.title} has charts or images. Value edits keep the workbook objects the library can save; they are not edited.",
                f"برگه {worksheet.title} نمودار یا تصویر دارد. ویرایش مقدارها اشیائی را که کتابخانه بتواند ذخیره کند نگه می‌دارد، اما آن‌ها را تغییر نمی‌دهد.",
            )
        )
    if formulas:
        warnings.append(
            warning(
                "formulas_not_evaluated",
                f"Sheet {worksheet.title} has formulas. Formulas are not evaluated or replaced with cached results.",
                f"برگه {worksheet.title} فرمول دارد. فرمول‌ها محاسبه نمی‌شوند و با نتیجه ذخیره‌شده جایگزین نمی‌شوند.",
            )
        )
    sheet = Sheet(
        name=worksheet.title,
        rows=rows,
        hidden=worksheet.sheet_state != "visible",
        merged_ranges=merged,
        chart_count=charts,
        image_count=images,
        comment_count=comments,
        hyperlink_count=hyperlinks,
        validation_count=validations,
        table_count=tables,
        hidden_row_count=hidden_rows,
        hidden_col_count=hidden_cols,
        declared_max_row=declared_row,
        declared_max_col=declared_col,
        actual_max_row=actual_max_row,
        actual_max_col=actual_max_col,
        formula_count=formulas,
    )
    return sheet, len(occupied)
