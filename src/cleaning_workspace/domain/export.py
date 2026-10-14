"""Write a transformed model to a new file. The source path is only read."""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import Cell as OpenpyxlCell

from cleaning_workspace.domain.models import Cell, WorkbookModel, row_is_fully_empty
from cleaning_workspace.domain.security import guard_formula_text
from cleaning_workspace.errors import AppError


def export_workbook(source: Path, model: WorkbookModel, destination: Path) -> None:
    if destination.resolve() == source.resolve():
        raise AppError(
            "Refusing to overwrite the source file.",
            "نوشتن روی فایل اصلی مجاز نیست.",
            code="source_protection",
            status_code=500,
        )
    destination.parent.mkdir(parents=True, exist_ok=True)
    if model.source_format == "csv":
        _export_csv(model, destination)
        return
    _export_xlsx(source, model, destination)


def _format_csv(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.hour == 0 and value.minute == 0 and value.second == 0:
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%dT%H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer() and abs(value) < 10**15:
        return str(int(value))
    if isinstance(value, str):
        return guard_formula_text(value)
    return str(value)


def _export_csv(model: WorkbookModel, destination: Path) -> None:
    sheet = model.sheets[0]
    width = sheet.width()
    with destination.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=model.delimiter or ",", lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
        for row in sheet.active_rows():
            exported = []
            for column in range(1, width + 1):
                cell = row.cells.get(column)
                exported.append(_format_csv(None if cell is None or cell.formula else cell.value))
            writer.writerow(exported)


def _rows_in_range(ref: str) -> set[int]:
    if ":" not in ref:
        return set()
    left, right = ref.split(":", 1)

    def row_number(token: str) -> int:
        digits = "".join(char for char in token if char.isdigit())
        return int(digits or 0)

    low, high = sorted((row_number(left), row_number(right)))
    if low < 1 or high < 1:
        return set()
    return set(range(low, high + 1))


def _assign(target: OpenpyxlCell, cell: Cell) -> None:
    if cell.formula:
        return
    target.value = cell.value
    if isinstance(cell.value, str) and str(cell.value).startswith("="):
        target.data_type = "s"
    if cell.number_format:
        target.number_format = cell.number_format


def _export_xlsx(source: Path, model: WorkbookModel, destination: Path) -> None:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(source.read_bytes()), data_only=False, keep_vba=model.has_vba)
    except Exception as exc:
        raise AppError(
            "The output workbook could not be prepared.",
            "آماده‌سازی فایل خروجی ناموفق بود.",
            code="export_failed",
            status_code=500,
        ) from exc
    for sheet in model.sheets:
        if sheet.name not in workbook.sheetnames:
            destination.unlink(missing_ok=True)
            raise AppError(
                "A sheet disappeared while exporting.",
                "یکی از برگه‌ها هنگام ساخت خروجی پیدا نشد.",
                code="export_failed",
                status_code=500,
            )
        worksheet = workbook[sheet.name]
        removed = {row.original_index for row in sheet.rows if row.removed}
        for merged in list(worksheet.merged_cells.ranges):
            rows = _rows_in_range(str(merged))
            if rows and rows <= removed:
                worksheet.unmerge_cells(str(merged))
        for row in sheet.rows:
            if row.removed:
                continue
            for column, cell in row.cells.items():
                if cell.formula:
                    continue
                changed = cell.value != cell.original_value or cell.number_format != cell.original_number_format
                new_column = column > sheet.actual_max_col
                if not changed and not new_column:
                    continue
                _assign(worksheet.cell(row.original_index, column), cell)
        for index in sorted(removed, reverse=True):
            worksheet.delete_rows(index, 1)
    try:
        buffer = io.BytesIO()
        workbook.save(buffer)
        destination.write_bytes(buffer.getvalue())
    except Exception as exc:
        destination.unlink(missing_ok=True)
        raise AppError(
            "The output workbook could not be saved.",
            "ذخیره فایل خروجی ناموفق بود.",
            code="export_failed",
            status_code=500,
        ) from exc


def count_fully_empty(model: WorkbookModel) -> dict[str, int]:
    return {sheet.name: sum(1 for row in sheet.active_rows() if row_is_fully_empty(row)) for sheet in model.sheets}
