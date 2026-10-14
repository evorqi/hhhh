import io
from datetime import datetime, timedelta, timezone

import openpyxl

from cleaning_workspace.api.app import create_app
from cleaning_workspace.config import Settings
from cleaning_workspace.domain.validate import sha256_file
from cleaning_workspace.services.jobs import JobService
from fastapi.testclient import TestClient
from tests.conftest import EXAMPLE


def _upload(client: TestClient, name: str, data: bytes, content_type: str = "application/octet-stream"):
    return client.post("/api/jobs", files={"file": (name, data, content_type)})


def test_upload_plan_execute_and_download_round_trip(client, service, tmp_path):
    from pathlib import Path
    sample = (Path(__file__).resolve().parents[1] / "web" / "examples" / "customers.csv").read_bytes()
    created = _upload(client, "customers.csv", sample, "application/vnd.ms-excel")
    assert created.status_code == 201, created.text
    job = created.json()
    assert job["detected_format"] == "csv"
    assert job["profile"]["fully_empty_rows"] >= 1
    source = service.files.source_path(job["id"])
    before = sha256_file(source)
    planned = client.post(f"/api/jobs/{job['id']}/plan", json={"request": EXAMPLE})
    assert planned.status_code == 200, planned.text
    plan = planned.json()["plan"]
    assert plan["interpreter"] == "rule_based"
    assert plan["blocking"] is False
    preview = next(op for op in plan["operations"] if op["type"] == "mark_exact_duplicates")["preview"]
    assert preview["affected_count"] >= 2
    executed = client.post(
        f"/api/jobs/{job['id']}/execute",
        json={"plan_id": plan["id"], "approve_destructive": False},
    )
    assert executed.status_code == 200, executed.text
    body = executed.json()
    assert body["status"] in {"completed", "completed_with_limitations"}
    assert body["result"]["validation"]["ok"] is True
    assert sha256_file(source) == before == body["source_sha256"]
    output = client.get(f"/api/jobs/{job['id']}/artifacts/output")
    assert output.status_code == 200
    assert output.content.startswith(b"\xef\xbb\xbf")
    text = output.content.decode("utf-8-sig")
    assert "علی رضایی" in text
    assert "'=1+1" in text
    assert "00123" in text
    assert "1,000" in text
    assert "01/02/2024" in text
    assert "1403/12/30" in text
    report = client.get(f"/api/jobs/{job['id']}/artifacts/report")
    assert report.status_code == 200
    assert report.json()["validation_passed"] is True
    again = client.post(
        f"/api/jobs/{job['id']}/execute",
        json={"plan_id": plan["id"], "approve_destructive": False},
    )
    assert again.status_code == 200
    assert again.json()["status"] == body["status"]


def test_destructive_plan_requires_approval_and_leaves_source(client, service):
    data = "نام,شهر\nعلی,تهران\nعلی,تهران\nسارا,یزد\n".encode()
    job = _upload(client, "people.csv", data).json()
    planned = client.post(f"/api/jobs/{job['id']}/plan", json={"request": "ردیف‌های تکراری را حذف کن"})
    assert planned.json()["status"] == "approval_required"
    plan_id = planned.json()["plan"]["id"]
    source = service.files.source_path(job["id"])
    before = source.read_bytes()
    denied = client.post(f"/api/jobs/{job['id']}/execute", json={"plan_id": plan_id, "approve_destructive": False})
    assert denied.status_code == 409
    assert source.read_bytes() == before
    assert client.get(f"/api/jobs/{job['id']}").json()["status"] == "approval_required"
    approved = client.post(f"/api/jobs/{job['id']}/execute", json={"plan_id": plan_id, "approve_destructive": True})
    assert approved.status_code == 200, approved.text
    assert approved.json()["result"]["validation"]["ok"] is True
    text = client.get(f"/api/jobs/{job['id']}/artifacts/output").content.decode("utf-8-sig")
    assert text.count("علی,تهران") == 1


def test_missing_customer_column_blocks_until_clarified(client):
    data = "عنوان,مقدار\nالف,1\nالف,1\n".encode()
    job = _upload(client, "t.csv", data).json()
    planned = client.post(
        f"/api/jobs/{job['id']}/plan",
        json={"request": "مشتری‌های تکراری را مشخص کن"},
    )
    assert planned.json()["status"] == "clarification_required"
    denied = client.post(
        f"/api/jobs/{job['id']}/execute",
        json={"plan_id": planned.json()["plan"]["id"], "approve_destructive": False},
    )
    assert denied.status_code == 409
    clarified = client.post(
        f"/api/jobs/{job['id']}/plan",
        json={"request": "مشتری‌های تکراری را مشخص کن", "clarifications": {"duplicate_key": "1"}},
    )
    assert clarified.json()["status"] == "plan_ready"
    assert clarified.json()["plan"]["operations"][0]["parameters"]["key_columns"] == [1]


def test_xls_misleading_mime_and_corrupt_xlsx_are_rejected(client):
    ole = bytes.fromhex("D0CF11E0A1B11AE1") + b"rest"
    rejected = _upload(client, "book.xlsx", ole, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert rejected.status_code == 415
    corrupt = _upload(client, "book.xlsx", b"not a workbook", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert corrupt.status_code == 422
    empty = _upload(client, "empty.csv", b"")
    assert empty.status_code == 422


def test_xlsx_named_as_csv_is_detected_by_signature(client):
    workbook = openpyxl.Workbook()
    workbook.active.append(["نام", "مبلغ"])
    workbook.active.append(["  علی  ", 10])
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = _upload(client, "data.csv", buffer.getvalue(), "text/csv")
    assert response.status_code == 201
    body = response.json()
    assert body["detected_format"] == "xlsx"
    assert any(item["code"] == "extension_mismatch" for item in body["inspection"]["warnings"])


def test_formula_is_not_replaced_and_text_equals_is_not_executed(client):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["نام", "محاسبه"])
    sheet.append(["  علی  ", "=1+1"])
    sheet["C2"] = "=CMD"
    sheet["C2"].data_type = "s"
    buffer = io.BytesIO()
    workbook.save(buffer)
    job = _upload(client, "calc.xlsx", buffer.getvalue()).json()
    planned = client.post(f"/api/jobs/{job['id']}/plan", json={"request": "فاصله‌های اضافی را بردار"}).json()
    executed = client.post(
        f"/api/jobs/{job['id']}/execute",
        json={"plan_id": planned["plan"]["id"], "approve_destructive": False},
    )
    assert executed.status_code == 200, executed.text
    output = client.get(f"/api/jobs/{job['id']}/artifacts/output").content
    loaded = openpyxl.load_workbook(io.BytesIO(output))
    assert loaded.active["A2"].value == "علی"
    assert loaded.active["B2"].value == "=1+1"
    assert loaded.active["B2"].data_type == "f"
    assert loaded.active["C2"].value == "=CMD"
    assert loaded.active["C2"].data_type == "s"


def test_unauthorized_token_and_unknown_job(tmp_path):
    settings = Settings(data_dir=tmp_path / "data", web_dir=tmp_path, access_token="secret-token")
    client = TestClient(create_app(settings, JobService(settings)))
    denied = client.get("/api/jobs/00000000-0000-0000-0000-000000000000")
    assert denied.status_code == 401
    assert "secret-token" not in denied.text
    ok = client.get("/api/meta", headers={"X-Workspace-Token": "secret-token"})
    assert ok.status_code == 200
    created = client.post(
        "/api/jobs",
        files={"file": ("a.csv", b"a,b\n1,2\n", "text/csv")},
        headers={"X-Workspace-Token": "secret-token"},
    )
    assert created.status_code == 201
    missing = client.get(
        "/api/jobs/00000000-0000-0000-0000-000000000000",
        headers={"X-Workspace-Token": "secret-token"},
    )
    assert missing.status_code == 404


def test_expired_job_deletes_files(client, service):
    job = _upload(client, "a.csv", "a,b\n1,2\n".encode()).json()
    service.repo.set_expires_at(job["id"], datetime.now(timezone.utc) - timedelta(hours=1))
    missing = client.get(f"/api/jobs/{job['id']}")
    assert missing.status_code == 404
    assert not service.files.source_path(job["id"]).exists()


def test_cancel_before_execute(client):
    job = _upload(client, "a.csv", "a\n x \n".encode()).json()
    cancelled = client.post(f"/api/jobs/{job['id']}/cancel")
    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "cancelled"
    planned = client.post(f"/api/jobs/{job['id']}/plan", json={"request": "فاصله را بردار"})
    assert planned.status_code == 409


def test_cp1256_round_trip(client):
    payload = "نام,مقدار\nمرحبا,1\n".encode("cp1256")
    created = _upload(client, "arabic.csv", payload)
    assert created.status_code == 201, created.text
    assert created.json()["inspection"]["encoding"] == "cp1256"
